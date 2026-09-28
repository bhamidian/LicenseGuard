from pathlib import Path
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

SKILL_SCRIPTS = Path(r"C:/Users/Abolfazl/.gapcode/plugins/cache/gapgpt/documents/26.722.10000/skills/documents/scripts")
sys.path.insert(0, str(SKILL_SCRIPTS))
from fa_docx import add_rtl_paragraph, set_font, set_rtl_paragraph, set_rtl_styles

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "docs" / "گزارش-پیشرفت-LicenseGuard.docx"
FONT = "Vazirmatn"
BLUE = "2E74B5"
MUTED = "667085"


def add_numbering(document):
    numbering = document.part.numbering_part.element
    abstract_ids = [int(node.get(qn("w:abstractNumId"))) for node in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(node.get(qn("w:numId"))) for node in numbering.findall(qn("w:num"))]
    abstract_id = max(abstract_ids, default=0) + 1
    num_id = max(num_ids, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    level.append(start)
    fmt = OxmlElement("w:numFmt")
    fmt.set(qn("w:val"), "bullet")
    level.append(fmt)
    text = OxmlElement("w:lvlText")
    text.set(qn("w:val"), "•")
    level.append(text)
    justify = OxmlElement("w:lvlJc")
    justify.set(qn("w:val"), "right")
    level.append(justify)
    ppr = OxmlElement("w:pPr")
    ind = OxmlElement("w:ind")
    ind.set(qn("w:right"), "720")
    ind.set(qn("w:hanging"), "360")
    ppr.append(ind)
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "720")
    tabs.append(tab)
    ppr.append(tabs)
    bidi = OxmlElement("w:bidi")
    ppr.append(bidi)
    level.append(ppr)
    rpr = OxmlElement("w:rPr")
    rfonts = OxmlElement("w:rFonts")
    rfonts.set(qn("w:ascii"), FONT)
    rfonts.set(qn("w:hAnsi"), FONT)
    rfonts.set(qn("w:cs"), FONT)
    rpr.append(rfonts)
    level.append(rpr)
    abstract.append(level)
    numbering.append(abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def add_bullet(document, text, num_id):
    paragraph = document.add_paragraph()
    ppr = paragraph._p.get_or_add_pPr()
    numpr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    numid = OxmlElement("w:numId")
    numid.set(qn("w:val"), str(num_id))
    numpr.append(ilvl)
    numpr.append(numid)
    ppr.append(numpr)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.line_spacing = 1.167
    run = paragraph.add_run(text)
    run.font.size = Pt(10.5)
    set_rtl_paragraph(paragraph, font=FONT, size=10.5)
    return paragraph


def add_body(document, text, bold_prefix=None):
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.10
    if bold_prefix and text.startswith(bold_prefix):
        run = paragraph.add_run(bold_prefix)
        run.bold = True
        paragraph.add_run(text[len(bold_prefix):])
    else:
        paragraph.add_run(text)
    set_rtl_paragraph(paragraph, font=FONT, size=11)
    return paragraph


def add_heading(document, text, level=1):
    paragraph = add_rtl_paragraph(document, text, style=f"Heading {level}", font=FONT)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    return paragraph


def main():
    doc = Document()
    set_rtl_styles(doc, font=FONT)
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor(35, 43, 55)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10
    for name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, "1F4D78", 8, 4),
    ]:
        style = doc.styles[name]
        style.font.name = FONT
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.add_run("LicenseGuard | گزارش وضعیت فنی")
    set_rtl_paragraph(header, font=FONT, size=8.5)
    for run in header.runs:
        run.font.color.rgb = RGBColor.from_string(MUTED)
        run.font.size = Pt(8.5)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("گزارش پیشرفت | ۲۵ سپتامبر ۲۰۲۶")
    set_rtl_paragraph(footer, font=FONT, size=8)
    for run in footer.runs:
        run.font.color.rgb = RGBColor.from_string(MUTED)
        run.font.size = Pt(8)

    title = add_rtl_paragraph(doc, "گزارش پیشرفت پروژه LicenseGuard", style="Title", font=FONT, size=23)
    title.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    title.paragraph_format.space_before = Pt(8)
    title.paragraph_format.space_after = Pt(4)
    title.runs[0].bold = True
    title.runs[0].font.color.rgb = RGBColor.from_string("0B2545")
    subtitle = add_rtl_paragraph(doc, "مدل دامنه، شماتیک پایگاه داده و وضعیت پیاده‌سازی", font=FONT, size=13)
    subtitle.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    subtitle.paragraph_format.space_after = Pt(14)
    subtitle.runs[0].font.color.rgb = RGBColor.from_string(MUTED)

    add_body(doc, "تاریخ گزارش: ۲۵ سپتامبر ۲۰۲۶ | بازه: نقطه‌ی وضعیت فعلی؛ تاریخ شروع بازه در مخزن ثبت نشده است.")
    lead = add_body(doc, "وضعیت کلی: زیرساخت اولیه‌ی مدل دامنه و شماتیک دیتابیس ایجاد شده است؛ قابلیت‌های اجرایی Backend هنوز پیاده‌سازی نشده‌اند.")
    lead.runs[0].bold = True

    num_id = add_numbering(doc)

    add_heading(doc, "خلاصه‌ی فعالیت‌های انجام‌شده")
    for item in [
        "نیازمندی‌های سند درباره‌ی Product، Plan، Feature، Subscription، License، Activation، تمدید و Audit مرور و به مدل دامنه نگاشت شد.",
        "روابط Customer، Product و Plan با Subscription تعریف شد؛ هر Subscription چند License تاریخی دارد و یک License جاری را مشخص می‌کند.",
        "نگاشت‌های EF Core در کلاس‌های جداگانه‌ی IEntityTypeConfiguration قرار گرفتند و DbContext با ApplyConfigurationsFromAssembly آن‌ها را بارگذاری می‌کند.",
        "Featureهای Product، Plan و License، Limitهای قابل‌گسترش، تاریخچه‌ی وضعیت و Audit Log به مدل دیتابیس افزوده شدند.",
        "برای تعلق License جاری به همان Subscription، قاعده‌ی Domain و Foreign Key مرکب در نگاشت EF تعریف شد.",
        "Migration اولیه‌ی InitialCreate برای MySQL تولید شد؛ هنوز روی دیتابیس اعمال نشده است.",
        "دو ERD شامل نمای کلی دامنه و نمای متمرکز بر چرخه‌ی License تهیه شد.",
    ]:
        add_bullet(doc, item, num_id)

    add_heading(doc, "وضعیت نسبت به زمان‌بندی")
    add_body(doc, "سند تسک ۲۰ روز کاری برای طراحی، Business Logic، API، Activation، Validation، تست و مستندسازی در نظر گرفته است. چون تاریخ شروع و نقاط زمان‌بندی پروژه ثبت نشده، درصد پیشرفت قابل محاسبه نیست. وضعیت فعلی مرحله‌ی آماده‌سازی مدل و طرح دیتابیس است و هنوز Backend قابل استفاده تحویل نشده است.")

    add_heading(doc, "کارهای تکمیل‌شده و باقی‌مانده")
    add_heading(doc, "تکمیل‌شده", 2)
    for item in [
        "Entityهای اصلی دامنه و روابط پایه‌ی Customer/Product/Plan/Subscription/License.",
        "پیکربندی مستقل EF Core برای موجودیت‌ها و تبدیل Value Objectها به ستون‌های دیتابیس.",
        "مدل واسط PlanFeature و رابطه‌ی License با Feature، Limit، Activation، تاریخچه و Audit.",
        "Migration اولیه و Model Snapshot، بدون اجرای آن روی دیتابیس.",
        "ERDهای قابل‌ویرایش Mermaid در docs/ERD.md و فایل‌های منبع docs/erd/.",
    ]:
        add_bullet(doc, item, num_id)

    add_heading(doc, "باقی‌مانده", 2)
    for item in [
        "لایه‌ی Application و APIهای مدیریتی و Client.",
        "چرخه‌ی اجرایی صدور، Activation/Deactivation، تمدید، Suspend/Resume، Revoke، Expire و Validation.",
        "کنترل رقابت درخواست‌ها برای سقف Activation و اعمال Featureها و Limitها در زمان Validation.",
        "احراز هویت، مجوزهای مدیریتی، Rate Limit، حفاظت از کلیدها و پالایش داده‌های حساس از Log.",
        "تست‌های واحد و یکپارچه، مستندات API و بررسی Migration روی نسخه‌ی واقعی MySQL/MariaDB.",
    ]:
        add_bullet(doc, item, num_id)

    add_heading(doc, "چالش‌ها و موانع فنی")
    for item in [
        "FK مرکب تعلق License جاری را به همان Subscription محدود می‌کند؛ انطباق Product اشتراک با Product پلن و سیاست تعویض License جاری هنوز باید در منطق برنامه enforce شود.",
        "وضعیت‌ها و تاریخچه مدل شده‌اند، اما انتقال وضعیت و ثبت تاریخچه باید در عملیات اتمیک پیاده و تست شود.",
        "Customer.AppUserId در جدول هست، اما FK آن در EF تعریف نشده است؛ کلاس AppUser فعلی نیز از IdentityRole ارث می‌برد و مدل حساب کاربری/نقش نیازمند بازبینی است.",
        "ابزار EF CLI نسخه‌ی 10.0.1 کنار Runtime نسخه‌ی 10.0.12 استفاده شد؛ Migration ساخته شد، اما هم‌نسخه‌کردن ابزارها پیشنهاد می‌شود.",
        "شماتیک migration هنوز روی MySQL/MariaDB اجرا نشده است؛ بنابراین رفتار نهایی Provider و قیود در دیتابیس واقعی تأیید نشده‌اند.",
    ]:
        add_bullet(doc, item, num_id)

    add_heading(doc, "نیازمندی‌ها و تصمیم‌های روشن‌شده")
    for item in [
        "Subscription به Customer، Product و Plan متصل است. رابطه‌ی License با Subscription یک‌به‌چند است تا کلیدهای قبلی حفظ شوند؛ یک License جاری برای هر Subscription در نظر گرفته شده است.",
        "Audit Log به License رابطه‌ی اختیاری دارد تا رخدادهای امنیتی بدون License قابل‌شناسایی نیز ثبت شوند.",
        "Featureها از طریق جدول‌های واسط ProductFeature، PlanFeature و LicenseFeature وصل می‌شوند. Limit با Code، Value و Unit تعریف شده تا قابل‌گسترش باشد.",
        "اتصال Billing/Payment و پنل مدیریتی هنوز پیاده‌سازی نشده و باید برای فاز اجرایی اولویت‌بندی شود.",
    ]:
        add_bullet(doc, item, num_id)

    add_heading(doc, "خروجی‌ها و بررسی")
    for item in [
        "ERD: docs/ERD.md؛ منابع جداگانه: docs/erd/01-domain-overview.mmd و docs/erd/02-license-lifecycle.mmd. جداول ASP.NET Identity برای خوانایی از ERDهای دامنه حذف شده‌اند.",
        "Migration: 02-Infrastructure/LicenseGuard.Infrastructure.EFCore/Persistence/Migrations/20260925200105_InitialCreate.cs و ApplicationDbContextModelSnapshot.cs.",
        "Build پروژه‌ی LicenseGuard.Infrastructure.EFCore موفق بود. تست خودکار اجرا نشد و Migration روی دیتابیس اعمال نشده است.",
    ]:
        add_bullet(doc, item, num_id)

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()

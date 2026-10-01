from pathlib import Path
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

SKILL_SCRIPTS = Path(r"C:/Users/Abolfazl/.gapcode/plugins/cache/gapgpt/documents/26.722.10000/skills/documents/scripts")
sys.path.insert(0, str(SKILL_SCRIPTS))
from fa_docx import add_rtl_paragraph, set_font, set_rtl_paragraph, set_rtl_styles

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "docs" / "گزارش-منطق-دامنه-LicenseGuard.docx"
FONT = "Vazirmatn"
BLUE = "2E74B5"
MUTED = "667085"
INK = "0B2545"


def add_numbering(document):
    numbering = document.part.numbering_part.element
    abstract_ids = [int(node.get(qn("w:abstractNumId"))) for node in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(node.get(qn("w:numId"))) for node in numbering.findall(qn("w:num"))]
    abstract_id = max(abstract_ids, default=0) + 1
    num_id = max(num_ids, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    level.append(start)
    fmt = OxmlElement("w:numFmt")
    fmt.set(qn("w:val"), "bullet")
    level.append(fmt)
    marker = OxmlElement("w:lvlText")
    marker.set(qn("w:val"), "•")
    level.append(marker)
    justify = OxmlElement("w:lvlJc")
    justify.set(qn("w:val"), "right")
    level.append(justify)
    ppr = OxmlElement("w:pPr")
    ind = OxmlElement("w:ind")
    ind.set(qn("w:right"), "720")
    ind.set(qn("w:hanging"), "360")
    ppr.append(ind)
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "720")
    tabs.append(tab)
    ppr.append(tabs)
    bidi = OxmlElement("w:bidi")
    ppr.append(bidi)
    level.append(ppr)
    rpr = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), FONT)
    fonts.set(qn("w:hAnsi"), FONT)
    fonts.set(qn("w:cs"), FONT)
    rpr.append(fonts)
    level.append(rpr)
    abstract.append(level)
    numbering.append(abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def add_bullet(document, text, num_id):
    paragraph = document.add_paragraph()
    ppr = paragraph._p.get_or_add_pPr()
    numpr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    numid = OxmlElement("w:numId")
    numid.set(qn("w:val"), str(num_id))
    numpr.append(ilvl)
    numpr.append(numid)
    ppr.append(numpr)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(5)
    paragraph.paragraph_format.line_spacing = 1.7
    paragraph.add_run(text)
    set_rtl_paragraph(paragraph, font=FONT, size=10.5)
    return paragraph


def add_body(document, text, bold_prefix=None):
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.7
    if bold_prefix and text.startswith(bold_prefix):
        lead = paragraph.add_run(bold_prefix)
        lead.bold = True
        paragraph.add_run(text[len(bold_prefix):])
    else:
        paragraph.add_run(text)
    set_rtl_paragraph(paragraph, font=FONT, size=10.5)
    return paragraph


def add_heading(document, text, level=1):
    paragraph = add_rtl_paragraph(document, text, style=f"Heading {level}", font=FONT, size=14 if level == 1 else 12)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    paragraph.paragraph_format.line_spacing = 1.7
    return paragraph


def main():
    doc = Document()
    set_rtl_styles(doc, font=FONT)
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(35, 43, 55)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.7
    for name, size, color, before, after in [
        ("Heading 1", 14, BLUE, 16, 8),
        ("Heading 2", 12, BLUE, 12, 6),
        ("Heading 3", 11, "1F4D78", 8, 4),
    ]:
        style = doc.styles[name]
        style.font.name = FONT
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.7
        style.paragraph_format.keep_with_next = True

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.add_run("LicenseGuard | گزارش فنی منطق دامنه")
    set_rtl_paragraph(header, font=FONT, size=8.5)
    for run in header.runs:
        run.font.color.rgb = RGBColor.from_string(MUTED)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("گزارش فنی | ۲۸ سپتامبر ۲۰۲۶")
    set_rtl_paragraph(footer, font=FONT, size=8)
    for run in footer.runs:
        run.font.color.rgb = RGBColor.from_string(MUTED)

    title = add_rtl_paragraph(doc, "گزارش تکمیل منطق دامنه و آزمون‌ها", style="Title", font=FONT, size=21)
    title.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    title.paragraph_format.space_before = Pt(8)
    title.paragraph_format.space_after = Pt(4)
    title.paragraph_format.line_spacing = 1.7
    title.runs[0].bold = True
    title.runs[0].font.color.rgb = RGBColor.from_string(INK)
    subtitle = add_rtl_paragraph(doc, "پیاده‌سازی قواعد Entityها، Value Objectها و چرخهٔ License", font=FONT, size=12)
    subtitle.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    subtitle.paragraph_format.space_after = Pt(10)
    subtitle.paragraph_format.line_spacing = 1.7
    subtitle.runs[0].font.color.rgb = RGBColor.from_string(MUTED)

    add_body(doc, "تاریخ گزارش: ۲۸ سپتامبر ۲۰۲۶ | دامنهٔ کار: منطق دامنه، خطاهای دامنه و آزمون‌های واحد.")
    lead = add_body(doc, "وضعیت کلی: رفتارهای اصلی مدل دامنه تکمیل شد؛ مجموعهٔ آزمون واحد با ۳۴ آزمون موفق و Build زیرساخت بدون هشدار و خطا اجرا شد.")
    lead.runs[0].bold = True

    num_id = add_numbering(doc)
    add_heading(doc, "فعالیت‌های انجام‌شده")
    for item in [
        "برای خطاهای قواعد کسب‌وکار، DomainException، DomainValidationException و DomainRuleViolationException اضافه شد تا خطای ورودی از نقض انتقال یا رابطهٔ مجاز تفکیک شود.",
        "در BaseEntity، شناسه و فیلدهای چرخهٔ عمر محصور شدند و عملیات SetCreatedBy، Deactivate، Reactivate، SoftDelete و Restore با قواعد سازگار اضافه شد.",
        "برای Product، Plan و Feature سازنده‌ها و متدهای تغییر اطلاعات و افزودن رابطه نوشته شد؛ اتصال Feature به Plan فقط وقتی پذیرفته می‌شود که Feature برای Product همان Plan تعریف شده باشد.",
        "Subscription اکنون بازهٔ زمانی، Customer/Product/Plan، لایسنس جاری، تغییر Plan، تمدید، لغو و انقضا را مدیریت می‌کند. تمدید ثبت‌شده به لایسنس جاری نیز اعمال می‌شود.",
        "License کلید تصادفی امن می‌سازد و فعال‌سازی، سقف تعداد، جلوگیری از Device/Instance تکراری، تعلیق، Resume، ابطال، انقضا، تمدید، Feature، Limit، امضا و تاریخچه را مدیریت می‌کند. اعتبارسنجی به Active بودن Subscription و Current License بودن همان License هم وابسته است.",
        "Activation عملیات Validation و Deactivation را با زمان UTC و وضعیت Active/Deactivated کنترل می‌کند. روابط درون‌حافظه‌ای License و Subscription نیز هنگام اتصال همگام می‌شوند.",
        "Value Objectهای شناسهٔ دستگاه و Instance، کلید و امضای License و تنظیم تمدید خودکار اعتبارسنجی شدند؛ کلید License فقط قالب ۶۴ رقم hexadecimal را می‌پذیرد.",
        "Entityهای Audit و تاریخچه به‌صورت سازنده‌محور پیاده شدند تا دادهٔ ثبت‌شده پس از ایجاد قابل تغییر آزادانه نباشد.",
    ]:
        add_bullet(doc, item, num_id)

    add_heading(doc, "آزمون‌ها و بررسی فنی")
    for item in [
        "پروژهٔ LicenseGuard.Domain.Tests شامل ۳۴ آزمون رفتاری برای Entityها، Value Objectها، enumها و LicenseSigningData است؛ تمام ۳۴ آزمون موفق شدند.",
        "Build پروژهٔ LicenseGuard.Infrastructure.EFCore موفق شد: صفر Warning و صفر Error.",
        "آزمون‌ها قواعد ورودی، روابط ناسازگار، انتقال وضعیت، تاریخ اعتبار، تمدید، سقف Activation، تکرار Machine/Instance، آزادشدن ظرفیت پس از Deactivation و پیوند Current License را بررسی می‌کنند.",
    ]:
        add_bullet(doc, item, num_id)

    add_heading(doc, "تصمیم‌های پیاده‌سازی و محدودهٔ باقی‌مانده")
    add_body(doc, "Setter عمومی برای تغییر دلخواه وضعیت‌ها در اختیار مصرف‌کننده قرار نگرفت؛ تغییرات از متدهای دامنه عبور می‌کنند تا اعتبارسنجی و تاریخچه هم‌زمان حفظ شود. کلید License در زمان ساخت با مولد رمزنگاری‌شده تولید می‌شود.")
    add_body(doc, "این تغییر، منطق داخل حافظهٔ مدل دامنه را پوشش می‌دهد. اجرای اتمیک در برابر درخواست‌های هم‌زمان، تراکنش و ذخیره‌سازی EF، API/Application، احراز هویت، تست یکپارچه با MySQL و پوشش خط به خط کد در این نوبت انجام نشده‌اند.")
    add_body(doc, "Migration جدیدی تولید یا اعمال نشد. بررسی EF Core وجود تغییر نسبت به آخرین Migration را گزارش کرد؛ چون پیش از این کار هم تغییرهای مدلِ ثبت‌نشده در workspace وجود داشت، منشأ اختلاف از این بررسی به‌تنهایی قابل انتساب نیست و باید جداگانه بازبینی شود.")

    doc.core_properties.title = "گزارش تکمیل منطق دامنه و آزمون‌های LicenseGuard"
    doc.core_properties.subject = "Domain behavior, exceptions and unit tests"
    doc.core_properties.author = "LicenseGuard"
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
